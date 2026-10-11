"""Writes a workbook model (from export-doc.mjs) to .xlsx. Preview writer; the app's
Dart writer follows the same rules: Arial, bold header with filter and frozen top row,
dates as real dates (dd/mm/yyyy), totals and averages as formulas, right-to-left for Arabic.

Usage: python3 write_xlsx.py model.json out.xlsx
"""
import json, sys
from datetime import date
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

INK, MUTED, ACCENT, ALERT = '15181C', '4F5861', '2152C4', '8C1D17'
HEAD_FILL = PatternFill('solid', fgColor='EEF2FB')
ALERT_FILL = PatternFill('solid', fgColor='FBECEA')
THIN = Side(style='thin', color='D5D9D4')
F = lambda **k: Font(name='Arial', size=k.pop('size', 10), **k)


def to_cell(value, typ):
    if value is None or value == '':
        return None
    if typ == 'date':
        y, m, d = map(int, value.split('-'))
        return date(y, m, d)
    return value


def write_table(ws, sheet, rtl):
    cols = sheet['columns']
    keys = [c['key'] for c in cols]
    letter = {k: get_column_letter(i + 1) for i, k in enumerate(keys)}
    for i, c in enumerate(cols, 1):
        cell = ws.cell(row=1, column=i, value=c['label'])
        cell.font = F(bold=True, color=INK)
        cell.fill = HEAD_FILL
        cell.border = Border(bottom=Side(style='thin', color=INK))
        cell.alignment = Alignment(horizontal='right' if (c['type'] in ('int', 'decimal')) != rtl else 'left', vertical='center', wrap_text=True)
        ws.column_dimensions[get_column_letter(i)].width = c.get('width', 14)
    row_formulas = sheet.get('rowFormulas') or {}
    for r, row in enumerate(sheet['rows'], 2):
        critical = row.get('style') == 'critical'
        for i, c in enumerate(cols, 1):
            k = c['key']
            if k in row_formulas:
                num, den = row_formulas[k]['ratio']
                value = f'=IF({letter[den]}{r}=0,"",ROUND({letter[num]}{r}/{letter[den]}{r},0))'
            else:
                value = to_cell(row.get(k), c['type'])
            cell = ws.cell(row=r, column=i, value=value)
            cell.font = F(color=ALERT if critical and k in ('issue', 'severity', 'alerts', 'days') else INK, bold=critical and k in ('issue', 'severity'))
            cell.border = Border(bottom=THIN)
            if c['type'] == 'date':
                cell.number_format = 'dd/mm/yyyy'
            elif c['type'] == 'int':
                cell.number_format = '0'
            elif c['type'] == 'decimal':
                cell.number_format = '0.0'
    last = len(sheet['rows']) + 1
    totals = sheet.get('totals')
    if totals:
        tr = last + 1
        ws.cell(row=tr, column=1, value=totals['label']).font = F(bold=True)
        for k in totals['sum']:
            ws.cell(row=tr, column=keys.index(k) + 1, value=f'=SUM({letter[k]}2:{letter[k]}{last})')
        for k, f in (totals.get('formulas') or {}).items():
            num, den = f['ratio']
            ws.cell(row=tr, column=keys.index(k) + 1, value=f'=IF({letter[den]}{tr}=0,"",ROUND({letter[num]}{tr}/{letter[den]}{tr},0))')
        for i in range(1, len(cols) + 1):
            cell = ws.cell(row=tr, column=i)
            cell.font = F(bold=True)
            cell.border = Border(top=Side(style='thin', color=INK))
    ws.freeze_panes = 'A2'
    if sheet['rows']:
        ws.auto_filter.ref = f'A1:{get_column_letter(len(cols))}{last}'


def write_summary(ws, sheet):
    for i, c in enumerate(sheet['columns'], 1):
        ws.column_dimensions[get_column_letter(i)].width = c.get('width', 14)
    r = 1
    for row in sheet['rows']:
        kind = row['kind']
        if kind == 'title':
            ws.cell(row=r, column=1, value=row['label']).font = F(size=14, bold=True, color=INK)
        elif kind == 'subtitle':
            ws.cell(row=r, column=1, value=row['label']).font = F(color=MUTED)
        elif kind == 'heading':
            c = ws.cell(row=r, column=1, value=row['label'])
            c.font = F(bold=True, color=ACCENT)
            for i in range(1, 4):
                ws.cell(row=r, column=i).border = Border(bottom=Side(style='thin', color=INK))
        elif kind == 'item':
            critical = row.get('style') == 'critical'
            ws.cell(row=r, column=1, value=row['label']).font = F(color=ALERT if critical else INK)
            if row.get('value') is not None:
                v = ws.cell(row=r, column=2, value=row['value'])
                v.font = F(bold=True, color=ALERT if critical else INK)
                v.number_format = '0'
            if row.get('note'):
                ws.cell(row=r, column=3, value=row['note']).font = F(color=MUTED)
            for i in range(1, 4):
                ws.cell(row=r, column=i).border = Border(bottom=THIN)
                if critical:
                    ws.cell(row=r, column=i).fill = ALERT_FILL
        r += 1


def main(model_path, out_path):
    model = json.load(open(model_path, encoding='utf-8'))
    rtl = model.get('dir') == 'rtl'
    wb = Workbook()
    wb.remove(wb.active)
    for sheet in model['sheets']:
        ws = wb.create_sheet(sheet['name'])
        ws.sheet_view.rightToLeft = rtl
        if sheet['kind'] == 'summary':
            write_summary(ws, sheet)
        else:
            write_table(ws, sheet, rtl)
        # Printing: one page wide, as many pages tall as needed; wide tables in landscape.
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.page_setup.paperSize = ws.PAPERSIZE_A4
        if sum(c.get('width', 14) for c in sheet['columns']) > 100:
            ws.page_setup.orientation = 'landscape'
        if sheet['kind'] == 'table':
            ws.print_title_rows = '1:1'
    if model.get('about'):
        ws = wb.create_sheet('About')
        ws.column_dimensions['A'].width = 22
        ws.column_dimensions['B'].width = 60
        lines = [('WO Tracker data copy', ''), ('Generated', model['about']['generatedAt']), ('Orders', model['about']['orders']),
                 ('Format version', model['about']['format']),
                 ('Note', 'A complete copy of every order. Columns marked "calculated" are for reading only.'),
                 ('Not in this file', model['about'].get('notIncluded', ''))]
        for i, (a, b) in enumerate(lines, 1):
            ws.cell(row=i, column=1, value=a).font = F(bold=True, size=14 if i == 1 else 10)
            ws.cell(row=i, column=2, value=b).font = F()
    wb.save(out_path)


if __name__ == '__main__':
    main(sys.argv[1], sys.argv[2])
